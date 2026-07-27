from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import Workbook

from document_reconciliation.adapters.excel import load_excel_records, records_from_matrix
from document_reconciliation.adapters.mineru import expand_html_table, extract_mineru_blocks, extract_mineru_records
from document_reconciliation.models import UNPARSEABLE


FIXTURES = Path(__file__).parent / "fixtures"


def _mineru_fixture() -> list:
    return json.loads((FIXTURES / "mineru_table_sample.json").read_text(encoding="utf-8"))


def test_expand_html_table_handles_rowspan_and_colspan() -> None:
    matrix = expand_html_table(_mineru_fixture()[0][0]["content"]["html"])
    assert matrix[:2] == [["Metric", "Forecast", "Forecast"], ["Metric", "2025A", "2026E"]]


def test_extracts_page_grouped_table_records() -> None:
    records = extract_mineru_records(_mineru_fixture())
    assert len(records) == 6
    assert {record.period for record in records} == {"2025A", "2026E"}


def test_mineru_trace_keeps_page_block_and_cell_locator() -> None:
    record = extract_mineru_records(_mineru_fixture())[0]
    assert record.source_trace["page_number"] == 1
    assert record.source_trace["block_index"] == 0
    assert record.source_trace["locator"] == "page:1:block:0:row:3:col:2"


def test_mineru_table_preserves_caption_and_bounded_footnote() -> None:
    record = extract_mineru_records(_mineru_fixture())[0]
    assert record.context == "Income statement (millions)"
    assert record.source_trace["caption_preview"] == "Income statement (millions)"
    assert record.source_trace["footnote_preview"] == "Illustrative values only"


def test_mineru_adapter_does_not_mutate_input() -> None:
    fixture = _mineru_fixture()
    before = deepcopy(fixture)
    extract_mineru_records(fixture)
    assert fixture == before


def test_non_table_paragraph_without_explicit_records_is_ignored() -> None:
    assert extract_mineru_records([[{"type": "paragraph", "text": "plain text"}]]) == []


def test_paragraph_explicit_records_are_supported() -> None:
    artifact = [[{"type": "paragraph", "content": {"records": [{"metric": "Items", "period": "2026E", "value": "2", "unit": "times"}]}}]]
    records = extract_mineru_records(artifact)
    assert records[0].metric_key == "items"
    assert records[0].normalized_value == "2"


def test_bad_table_html_returns_no_records() -> None:
    assert extract_mineru_records([[{"type": "table", "content": {"html": ""}}]]) == []


def test_extract_mineru_blocks_uses_metadata_not_html_dump() -> None:
    blocks = extract_mineru_blocks(_mineru_fixture())
    assert blocks[0]["bbox"] == [10, 20, 300, 180]
    assert "html" not in blocks[0]


def test_matrix_to_records_marks_invalid_numbers_unparseable() -> None:
    matrix = [["Metric", "2026E"], ["Revenue", "unknown"]]
    records = records_from_matrix(matrix, source="right", sheet_name="Input")
    assert records[0].parse_status == UNPARSEABLE


def test_excel_adapter_reads_selected_sheet_and_source_trace(tmp_path: Path) -> None:
    path = tmp_path / "anonymous.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Income"
    sheet.append(["Context", "Metric", "Unit", "2025A", "2026E"])
    sheet.append(["income", "Revenue", "million", 100, 120])
    other = workbook.create_sheet("Other")
    other.append(["Metric", "2025A"])
    other.append(["Ignore", 1])
    workbook.save(path)
    workbook.close()
    records = load_excel_records(path, sheets=["Income"])
    assert len(records) == 2
    assert records[0].source_trace["sheet"] == "Income"
    assert records[0].normalized_unit == "million"


def test_excel_adapter_rejects_missing_selected_sheet(tmp_path: Path) -> None:
    path = tmp_path / "anonymous.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()
    with pytest.raises(ValueError, match="requested sheets"):
        load_excel_records(path, sheets=["Missing"])
