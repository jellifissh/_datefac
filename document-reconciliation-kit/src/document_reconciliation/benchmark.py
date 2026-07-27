"""Reusable, VERIFIED-only review-pack benchmark metrics."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook

from .normalization import normalize_numeric
from .reconciliation import MATCH


REVIEW_PACK_COLUMNS = (
    "review_id",
    "context",
    "metric_key",
    "metric_display_name",
    "period",
    "unit",
    "left_value",
    "right_value",
    "comparison_status",
    "ground_truth_value",
    "ground_truth_unit",
    "ground_truth_review_status",
    "left_correct",
    "right_correct",
    "comparison_correct",
    "error_severity",
    "evidence_note",
    "reviewer_note",
)


def build_review_pack_rows(comparison_rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, comparison in enumerate(comparison_rows, start=1):
        row = dict(comparison)
        rows.append(
            {
                "review_id": f"review-{index:04d}",
                "context": row.get("context", ""),
                "metric_key": row.get("metric_key", ""),
                "metric_display_name": row.get("metric_display_name", ""),
                "period": row.get("period", ""),
                "unit": row.get("left_unit") or row.get("right_unit") or "",
                "left_value": row.get("left_value", ""),
                "right_value": row.get("right_value", ""),
                "comparison_status": row.get("status", ""),
                "ground_truth_value": "",
                "ground_truth_unit": "",
                "ground_truth_review_status": "",
                "left_correct": "",
                "right_correct": "",
                "comparison_correct": "",
                "error_severity": "",
                "evidence_note": "",
                "reviewer_note": "",
            }
        )
    return rows


def write_review_pack(path: str | Path, rows: Sequence[Mapping[str, Any]]) -> Path:
    output_path = Path(path)
    if output_path.exists():
        raise ValueError(f"review pack already exists: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "review_pack"
    worksheet.append(list(REVIEW_PACK_COLUMNS))
    for row in rows:
        worksheet.append([_excel_value(row.get(column, "")) for column in REVIEW_PACK_COLUMNS])
    worksheet.freeze_panes = "A2"
    workbook.save(output_path)
    workbook.close()
    return output_path


def read_review_pack(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        header = tuple("" if value is None else str(value) for value in next(iterator, ()))
        validate_review_pack_schema(header)
        return [
            {column: "" if value is None else value for column, value in zip(REVIEW_PACK_COLUMNS, values, strict=False)}
            for values in iterator
        ]
    finally:
        workbook.close()


def validate_review_pack_schema(header: Sequence[str]) -> None:
    if tuple(header) != REVIEW_PACK_COLUMNS:
        raise ValueError("review pack schema mismatch")


def calculate_benchmark_metrics(rows: Iterable[Mapping[str, Any]]) -> dict[str, Any]:
    all_rows = [dict(row) for row in rows]
    verified_rows = [row for row in all_rows if str(row.get("ground_truth_review_status", "")).upper() == "VERIFIED"]
    metrics: dict[str, Any] = {
        "sampled_cell_count": len(all_rows),
        "verified_cell_count": len(verified_rows),
        "left_error_count": 0,
        "right_error_count": 0,
        "both_wrong_same_value_count": 0,
        "both_wrong_different_value_count": 0,
        "comparison_true_positive_count": 0,
        "comparison_false_positive_count": 0,
        "comparison_false_negative_count": 0,
        "comparison_true_negative_count": 0,
        "precision": 0.0,
        "recall": 0.0,
        "false_positive_rate": 0.0,
    }
    for row in verified_rows:
        left_correct = resolve_correctness(row, system="left")
        right_correct = resolve_correctness(row, system="right")
        any_error = not left_correct or not right_correct
        flagged = str(row.get("comparison_status", "")).upper() != MATCH
        if not left_correct:
            metrics["left_error_count"] += 1
        if not right_correct:
            metrics["right_error_count"] += 1
        if not left_correct and not right_correct:
            if normalize_benchmark_value(row.get("left_value")) == normalize_benchmark_value(row.get("right_value")):
                metrics["both_wrong_same_value_count"] += 1
            else:
                metrics["both_wrong_different_value_count"] += 1
        if flagged and any_error:
            metrics["comparison_true_positive_count"] += 1
        elif flagged:
            metrics["comparison_false_positive_count"] += 1
        elif any_error:
            metrics["comparison_false_negative_count"] += 1
        else:
            metrics["comparison_true_negative_count"] += 1
    metrics["precision"] = _safe_ratio(metrics["comparison_true_positive_count"], metrics["comparison_true_positive_count"] + metrics["comparison_false_positive_count"])
    metrics["recall"] = _safe_ratio(metrics["comparison_true_positive_count"], metrics["comparison_true_positive_count"] + metrics["comparison_false_negative_count"])
    metrics["false_positive_rate"] = _safe_ratio(metrics["comparison_false_positive_count"], metrics["comparison_false_positive_count"] + metrics["comparison_true_negative_count"])
    return metrics


def resolve_correctness(row: Mapping[str, Any], *, system: str) -> bool:
    explicit = _parse_bool(row.get(f"{system}_correct"))
    if explicit is not None:
        return explicit
    truth = normalize_benchmark_value(row.get("ground_truth_value"))
    value = normalize_benchmark_value(row.get(f"{system}_value"))
    if not truth:
        return False
    truth_unit = str(row.get("ground_truth_unit", "")).strip()
    unit = str(row.get("unit", "")).strip()
    return value == truth and (not truth_unit or not unit or unit == truth_unit)


def render_benchmark_markdown(metrics: Mapping[str, Any]) -> str:
    lines = ["# Reconciliation benchmark summary", ""]
    for key in sorted(metrics):
        lines.append(f"- {key}: {metrics[key]}")
    return "\n".join(lines) + "\n"


def write_benchmark_summary(rows: Iterable[Mapping[str, Any]], output_dir: str | Path) -> dict[str, str]:
    output_path = _safe_empty_output_dir(output_dir)
    metrics = calculate_benchmark_metrics(rows)
    json_path = output_path / "benchmark_summary.json"
    markdown_path = output_path / "benchmark_summary.md"
    import json

    json_path.write_text(json.dumps(metrics, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    markdown_path.write_text(render_benchmark_markdown(metrics), encoding="utf-8")
    return {"json_path": str(json_path), "markdown_path": str(markdown_path)}


def normalize_benchmark_value(value: Any) -> str:
    return normalize_numeric(value) or ""


def _parse_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"true", "yes", "1"}:
        return True
    if normalized in {"false", "no", "0"}:
        return False
    return None


def _safe_ratio(numerator: int, denominator: int) -> float:
    return 0.0 if denominator == 0 else numerator / denominator


def _excel_value(value: Any) -> Any:
    return "" if value is None else value


def _safe_empty_output_dir(output_dir: str | Path) -> Path:
    path = Path(output_dir)
    if path.exists() and not path.is_dir():
        raise ValueError(f"output path is not a directory: {path}")
    if path.exists() and any(path.iterdir()):
        raise ValueError(f"output directory is not empty: {path}")
    path.mkdir(parents=True, exist_ok=True)
    return path
