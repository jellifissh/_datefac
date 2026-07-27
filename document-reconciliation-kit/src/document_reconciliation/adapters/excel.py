"""Read selected Excel worksheets into normalized records."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable, Iterable, Sequence

from openpyxl import load_workbook

from ..models import NormalizedRecord, UNPARSEABLE
from ..normalization import bounded_preview, compact_text, normalize_metric_key, normalize_numeric, normalize_period, normalize_unit


def load_excel_records(
    path: str | Path,
    *,
    sheets: Sequence[str] | None = None,
    source: str = "right",
    metric_normalizer: Callable[[Any], str] = normalize_metric_key,
) -> list[NormalizedRecord]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        selected = list(sheets) if sheets is not None else list(workbook.sheetnames)
        missing = [name for name in selected if name not in workbook.sheetnames]
        if missing:
            raise ValueError(f"requested sheets are missing: {', '.join(missing)}")
        records: list[NormalizedRecord] = []
        for sheet_name in selected:
            matrix = [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
            records.extend(records_from_matrix(matrix, source=source, sheet_name=sheet_name, metric_normalizer=metric_normalizer))
        return records
    finally:
        workbook.close()


def records_from_matrix(
    matrix: Iterable[Sequence[Any]],
    *,
    source: str,
    sheet_name: str,
    metric_normalizer: Callable[[Any], str] = normalize_metric_key,
) -> list[NormalizedRecord]:
    rows = [list(row) for row in matrix]
    header_index, period_columns = _find_header(rows)
    if header_index is None:
        return []
    header = rows[header_index]
    metric_column = _column_named(header, {"metric", "item", "label", "指标", "项目", "科目"})
    context_column = _column_named(header, {"context", "statement", "section", "context_name", "类别"})
    unit_column = _column_named(header, {"unit", "单位"})
    if metric_column is None:
        period_indexes = {index for index, _ in period_columns}
        metric_column = next((index for index in range(len(header)) if index not in period_indexes), 0)
    records: list[NormalizedRecord] = []
    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 1):
        metric_display_name = _cell(row, metric_column)
        if not metric_display_name:
            continue
        context = _cell(row, context_column) if context_column is not None else compact_text(sheet_name)
        unit = normalize_unit(_cell(row, unit_column)) if unit_column is not None else None
        for column_index, period in period_columns:
            raw_value = _value_at(row, column_index)
            if raw_value is None or compact_text(raw_value) == "":
                continue
            normalized_value = normalize_numeric(raw_value)
            trace = {
                "source": source,
                "sheet": sheet_name,
                "row": row_index + 1,
                "column": column_index + 1,
                "locator": f"sheet:{sheet_name}:row:{row_index + 1}:col:{column_index + 1}",
            }
            records.append(
                NormalizedRecord(
                    source=source,
                    context=context,
                    metric_key=metric_normalizer(metric_display_name),
                    metric_display_name=metric_display_name,
                    period=period,
                    normalized_value=normalized_value,
                    normalized_unit=unit,
                    evidence_preview=bounded_preview(f"{sheet_name} | {metric_display_name} | {period} | {raw_value}"),
                    source_trace=trace,
                    parse_status="PARSED" if normalized_value is not None else UNPARSEABLE,
                )
            )
    return records


def _find_header(rows: list[list[Any]]) -> tuple[int | None, list[tuple[int, str]]]:
    for row_index, row in enumerate(rows[:20]):
        periods = [(column_index, period) for column_index, value in enumerate(row) if (period := normalize_period(value))]
        if periods:
            return row_index, periods
    return None, []


def _column_named(header: Sequence[Any], names: set[str]) -> int | None:
    normalized = {compact_text(name).casefold() for name in names}
    return next((index for index, value in enumerate(header) if compact_text(value).casefold() in normalized), None)


def _cell(row: Sequence[Any], index: int) -> str:
    return compact_text(_value_at(row, index))


def _value_at(row: Sequence[Any], index: int) -> Any:
    return row[index] if 0 <= index < len(row) else None
