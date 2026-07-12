"""R7CF project-necessity benchmark pilot.

The pilot inventories local report packages and builds/evaluates a one-report
ground-truth review pack. It does not parse PDFs, run MinerU/OCR/LLM/VLM, write
clean_data, or open readiness gates.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Sequence
import hashlib
import json
import os
import re

from openpyxl import Workbook, load_workbook

from datefac_agent.reconciliation.real_artifact_compatibility_348n import (
    MATCH,
    READINESS_GATES_CLOSED,
    load_and_reconcile_real_artifacts,
    normalize_numeric_value,
)

TASK_ID = "348N-R7CF"
BENCHMARK_VERSION = "r7cf_project_necessity_ground_truth_benchmark_pilot_v1"
ANJING_REPORT_ID = "H3_AP202606081823352906_1"

DEFAULT_INVENTORY_ROOTS: tuple[str, ...] = (
    r"D:\_datefac_agent",
    r"D:\_datefac",
    r"E:\mineru_lab",
    r"E:\_datefac_toolbench",
)

REVIEW_PACK_COLUMNS: tuple[str, ...] = (
    "report_id",
    "pdf_basename",
    "pdf_page",
    "pdf_locator_or_bbox",
    "statement_context",
    "metric",
    "period",
    "unit",
    "mineru_value",
    "original_value",
    "datefac_status",
    "pdf_ground_truth_value",
    "pdf_ground_truth_unit",
    "ground_truth_review_status",
    "mineru_correct",
    "original_correct",
    "datefac_decision_correct",
    "error_severity",
    "evidence_note",
    "reviewer_note",
)

READY = "READY"
MISSING_PDF = "MISSING_PDF"
MISSING_MINERU_ARTIFACT = "MISSING_MINERU_ARTIFACT"
MISSING_ORIGINAL_ARTIFACT = "MISSING_ORIGINAL_ARTIFACT"
AMBIGUOUS_PAIRING = "AMBIGUOUS_PAIRING"

BENCHMARK_METHOD_VALID = "BENCHMARK_METHOD_VALID"
BENCHMARK_METHOD_INVALID = "BENCHMARK_METHOD_INVALID"
NEEDS_MORE_VERIFIED_CELLS = "NEEDS_MORE_VERIFIED_CELLS"
ONE_REPORT_SUGGESTS_LIGHTWEIGHT_QA_VALUE = "ONE_REPORT_SUGGESTS_LIGHTWEIGHT_QA_VALUE"
ONE_REPORT_SUGGESTS_LOW_VALUE = "ONE_REPORT_SUGGESTS_LOW_VALUE"

KNOWN_ORIGINAL_ID_HINTS: dict[str, str] = {
    "datefac_raw_material_anjing_foods.xlsx": ANJING_REPORT_ID,
}

CANONICAL_BENCHMARK_PATHS: dict[str, dict[str, tuple[str, ...]]] = {
    ANJING_REPORT_ID: {
        "pdf": (
            r"E:\mineru_lab\input\H3_AP202606081823352906_1.pdf",
        ),
        "mineru": (
            r"E:\mineru_lab\output_new\H3_AP202606081823352906_1\auto\H3_AP202606081823352906_1_content_list_v2.json",
        ),
        "original": (
            r"D:\_datefac_agent\output\datefac_raw_material_anjing_foods.xlsx",
        ),
    },
}

REPORT_ID_RE = re.compile(r"H3_AP\d+_\d+", re.IGNORECASE)
JSON_INDENT = 2

HIGH_VALUE_METRIC_PRIORITIES: tuple[str, ...] = (
    "revenue",
    "parent_net_profit",
    "net_profit",
    "eps_diluted",
    "eps",
    "roe",
    "pe",
    "pb",
    "gross_margin",
    "net_margin",
    "total_assets",
    "total_liabilities",
    "equity",
    "operating_cash_flow",
    "investing_cash_flow",
    "financing_cash_flow",
    "debt_to_asset_ratio",
    "total_asset_turnover",
)

PERIOD_SUFFIX_PRIORITY: dict[str, int] = {
    "E": 0,
    "F": 1,
    "A": 2,
    "": 3,
}

FORBIDDEN_COMMITTED_SUFFIXES = frozenset({".pdf", ".xlsx", ".xlsm", ".xls"})


@dataclass(frozen=True)
class ReportPackageCandidate:
    report_id: str
    status: str
    pdf_paths: tuple[str, ...]
    mineru_artifact_paths: tuple[str, ...]
    original_artifact_paths: tuple[str, ...]
    selected_pdf_path: str | None = None
    selected_mineru_artifact_path: str | None = None
    selected_original_artifact_path: str | None = None
    reason: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {
            "report_id": self.report_id,
            "status": self.status,
            "pdf_paths": list(self.pdf_paths),
            "mineru_artifact_paths": list(self.mineru_artifact_paths),
            "original_artifact_paths": list(self.original_artifact_paths),
            "selected_pdf_path": self.selected_pdf_path,
            "selected_mineru_artifact_path": self.selected_mineru_artifact_path,
            "selected_original_artifact_path": self.selected_original_artifact_path,
            "reason": self.reason,
        }


def inventory_report_packages(roots: Sequence[str | Path]) -> list[ReportPackageCandidate]:
    buckets: dict[str, dict[str, list[str]]] = {}
    for root in roots:
        root_path = Path(root)
        if not root_path.exists():
            continue
        for file_path in _iter_inventory_files(root_path):
            report_id = extract_report_id(file_path)
            if not report_id:
                continue
            file_kind = classify_inventory_file(file_path)
            if not file_kind:
                continue
            bucket = buckets.setdefault(report_id, {"pdf": [], "mineru": [], "original": []})
            bucket[file_kind].append(str(file_path))

    candidates = [
        classify_report_package(
            report_id,
            pdf_paths=_dedupe_paths(bucket["pdf"]),
            mineru_paths=_dedupe_paths(bucket["mineru"]),
            original_paths=_dedupe_paths(bucket["original"]),
        )
        for report_id, bucket in buckets.items()
    ]
    return sorted(candidates, key=lambda candidate: (candidate.status != READY, candidate.report_id.lower()))


def classify_report_package(
    report_id: str,
    *,
    pdf_paths: Sequence[str | Path],
    mineru_paths: Sequence[str | Path],
    original_paths: Sequence[str | Path],
) -> ReportPackageCandidate:
    pdf_values = _dedupe_paths(pdf_paths)
    mineru_values = _dedupe_paths(mineru_paths)
    original_values = _dedupe_paths(original_paths)
    if not pdf_values:
        return ReportPackageCandidate(
            report_id=report_id,
            status=MISSING_PDF,
            pdf_paths=pdf_values,
            mineru_artifact_paths=mineru_values,
            original_artifact_paths=original_values,
            reason="no original PDF with stable report identity was found",
        )
    if not mineru_values:
        return ReportPackageCandidate(
            report_id=report_id,
            status=MISSING_MINERU_ARTIFACT,
            pdf_paths=pdf_values,
            mineru_artifact_paths=mineru_values,
            original_artifact_paths=original_values,
            reason="no content_list_v2 MinerU artifact with stable report identity was found",
        )
    if not original_values:
        return ReportPackageCandidate(
            report_id=report_id,
            status=MISSING_ORIGINAL_ARTIFACT,
            pdf_paths=pdf_values,
            mineru_artifact_paths=mineru_values,
            original_artifact_paths=original_values,
            reason="no original Excel/JSON artifact with stable report identity was found",
        )
    if len(pdf_values) != 1 or len(mineru_values) != 1 or len(original_values) != 1:
        return ReportPackageCandidate(
            report_id=report_id,
            status=AMBIGUOUS_PAIRING,
            pdf_paths=pdf_values,
            mineru_artifact_paths=mineru_values,
            original_artifact_paths=original_values,
            reason="one or more package roles have multiple possible files",
        )
    return ReportPackageCandidate(
        report_id=report_id,
        status=READY,
        pdf_paths=pdf_values,
        mineru_artifact_paths=mineru_values,
        original_artifact_paths=original_values,
        selected_pdf_path=pdf_values[0],
        selected_mineru_artifact_path=mineru_values[0],
        selected_original_artifact_path=original_values[0],
        reason="unique PDF, MinerU content_list_v2 artifact, and original artifact were paired",
    )


def select_one_report_package(candidates: Sequence[ReportPackageCandidate]) -> ReportPackageCandidate | None:
    canonical = resolve_canonical_report_package(candidates)
    if canonical is not None:
        return canonical
    ready_candidates = [candidate for candidate in candidates if candidate.status == READY]
    for candidate in ready_candidates:
        if candidate.report_id == ANJING_REPORT_ID:
            return candidate
    return ready_candidates[0] if ready_candidates else None


def resolve_canonical_report_package(
    candidates: Sequence[ReportPackageCandidate],
) -> ReportPackageCandidate | None:
    by_report_id = {candidate.report_id: candidate for candidate in candidates}
    candidate = by_report_id.get(ANJING_REPORT_ID)
    if candidate is None:
        return None
    canonical_paths = CANONICAL_BENCHMARK_PATHS.get(ANJING_REPORT_ID, {})
    selected_pdf = _select_preferred_path(candidate.pdf_paths, canonical_paths.get("pdf", ()))
    selected_mineru = _select_preferred_path(candidate.mineru_artifact_paths, canonical_paths.get("mineru", ()))
    selected_original = _select_preferred_path(candidate.original_artifact_paths, canonical_paths.get("original", ()))
    if not selected_pdf or not selected_mineru or not selected_original:
        return None
    return ReportPackageCandidate(
        report_id=candidate.report_id,
        status=READY,
        pdf_paths=(selected_pdf,),
        mineru_artifact_paths=(selected_mineru,),
        original_artifact_paths=(selected_original,),
        selected_pdf_path=selected_pdf,
        selected_mineru_artifact_path=selected_mineru,
        selected_original_artifact_path=selected_original,
        reason="canonical benchmark trio selected from inventory",
    )


def build_one_report_review_pack_rows(
    package: ReportPackageCandidate,
    *,
    sample_size: int = 30,
) -> list[dict[str, Any]]:
    if package.status != READY:
        raise ValueError(f"package {package.report_id} is not READY")
    if not package.selected_mineru_artifact_path or not package.selected_original_artifact_path:
        raise ValueError(f"package {package.report_id} has no selected artifacts")
    reconciliation = load_and_reconcile_real_artifacts(
        package.selected_mineru_artifact_path,
        package.selected_original_artifact_path,
    )
    sampled_rows = sample_high_value_cells(reconciliation["comparison_rows"], sample_size=sample_size)
    pdf_basename = Path(package.selected_pdf_path or "").name
    return [
        build_review_pack_row(package.report_id, pdf_basename, comparison_row)
        for comparison_row in sampled_rows
    ]


def build_review_pack_row(
    report_id: str,
    pdf_basename: str,
    comparison_row: dict[str, Any],
) -> dict[str, Any]:
    mineru_trace = _safe_mapping(comparison_row.get("source_trace", {}).get("mineru"))
    locator = _trace_locator_or_bbox(mineru_trace)
    return {
        "report_id": report_id,
        "pdf_basename": pdf_basename,
        "pdf_page": mineru_trace.get("page_number", ""),
        "pdf_locator_or_bbox": locator,
        "statement_context": comparison_row.get("statement_context", ""),
        "metric": comparison_row.get("metric", ""),
        "period": comparison_row.get("period", ""),
        "unit": comparison_row.get("mineru_unit") or comparison_row.get("original_unit") or "",
        "mineru_value": comparison_row.get("mineru_value", ""),
        "original_value": comparison_row.get("original_value", ""),
        "datefac_status": comparison_row.get("status", ""),
        "pdf_ground_truth_value": "",
        "pdf_ground_truth_unit": "",
        "ground_truth_review_status": "",
        "mineru_correct": "",
        "original_correct": "",
        "datefac_decision_correct": "",
        "error_severity": "",
        "evidence_note": "",
        "reviewer_note": "",
    }


def sample_high_value_cells(comparison_rows: Sequence[dict[str, Any]], *, sample_size: int = 30) -> list[dict[str, Any]]:
    if sample_size <= 0:
        return []
    unique_rows: dict[tuple[str, str, str], dict[str, Any]] = {}
    for row in comparison_rows:
        key = (
            str(row.get("statement_context", "")),
            str(row.get("metric_key", row.get("metric", ""))),
            str(row.get("period", "")),
        )
        unique_rows.setdefault(key, dict(row))
    ordered = sorted(unique_rows.values(), key=_sample_sort_key)
    return ordered[:sample_size]


def write_review_pack(path: str | Path, rows: Sequence[dict[str, Any]]) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ground_truth_review"
    worksheet.append(list(REVIEW_PACK_COLUMNS))
    for row in rows:
        worksheet.append([_excel_safe(row.get(column, "")) for column in REVIEW_PACK_COLUMNS])
    worksheet.freeze_panes = "A2"
    workbook.save(output_path)
    workbook.close()
    return output_path


def read_review_pack(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows_iter = worksheet.iter_rows(values_only=True)
        header = tuple(str(value or "") for value in next(rows_iter, ()))
        validate_review_pack_schema(header)
        rows: list[dict[str, Any]] = []
        for values in rows_iter:
            row = {
                column: "" if value is None else value
                for column, value in zip(REVIEW_PACK_COLUMNS, values, strict=False)
            }
            rows.append(row)
        return rows
    finally:
        workbook.close()


def validate_review_pack_schema(header: Sequence[str]) -> None:
    if tuple(header) != REVIEW_PACK_COLUMNS:
        raise ValueError("review pack schema mismatch")


def calculate_benchmark_metrics(rows: Sequence[dict[str, Any]]) -> dict[str, Any]:
    verified_rows = [row for row in rows if _is_verified(row.get("ground_truth_review_status"))]
    metric_state = {
        "sampled_cell_count": len(rows),
        "verified_cell_count": len(verified_rows),
        "mineru_error_count": 0,
        "original_error_count": 0,
        "both_wrong_same_value_count": 0,
        "both_wrong_different_value_count": 0,
        "datefac_true_positive_count": 0,
        "datefac_false_positive_count": 0,
        "datefac_false_negative_count": 0,
        "datefac_true_negative_count": 0,
        "precision": 0.0,
        "recall": 0.0,
        "false_positive_rate": 0.0,
        "high_severity_error_count": 0,
    }
    for row in verified_rows:
        mineru_correct = resolve_correctness(row, system="mineru")
        original_correct = resolve_correctness(row, system="original")
        mineru_wrong = not mineru_correct
        original_wrong = not original_correct
        any_error = mineru_wrong or original_wrong
        datefac_review_flag = str(row.get("datefac_status", "")).upper() != MATCH

        if mineru_wrong:
            metric_state["mineru_error_count"] += 1
        if original_wrong:
            metric_state["original_error_count"] += 1
        if mineru_wrong and original_wrong:
            mineru_value = normalize_benchmark_value(row.get("mineru_value"))
            original_value = normalize_benchmark_value(row.get("original_value"))
            if mineru_value == original_value:
                metric_state["both_wrong_same_value_count"] += 1
            else:
                metric_state["both_wrong_different_value_count"] += 1
        if any_error and _is_high_severity(row.get("error_severity")):
            metric_state["high_severity_error_count"] += 1

        if datefac_review_flag and any_error:
            metric_state["datefac_true_positive_count"] += 1
        elif datefac_review_flag and not any_error:
            metric_state["datefac_false_positive_count"] += 1
        elif not datefac_review_flag and any_error:
            metric_state["datefac_false_negative_count"] += 1
        else:
            metric_state["datefac_true_negative_count"] += 1

    metric_state["precision"] = _safe_ratio(
        metric_state["datefac_true_positive_count"],
        metric_state["datefac_true_positive_count"] + metric_state["datefac_false_positive_count"],
    )
    metric_state["recall"] = _safe_ratio(
        metric_state["datefac_true_positive_count"],
        metric_state["datefac_true_positive_count"] + metric_state["datefac_false_negative_count"],
    )
    metric_state["false_positive_rate"] = _safe_ratio(
        metric_state["datefac_false_positive_count"],
        metric_state["datefac_false_positive_count"] + metric_state["datefac_true_negative_count"],
    )
    return metric_state


def build_benchmark_summary(
    rows: Sequence[dict[str, Any]],
    *,
    candidate_report_package_count: int,
    ready_report_package_count: int,
) -> dict[str, Any]:
    metrics = calculate_benchmark_metrics(rows)
    decision = decide_provisional_project_value(metrics, ready_report_package_count=ready_report_package_count)
    return {
        "task_id": TASK_ID,
        "benchmark_version": BENCHMARK_VERSION,
        "candidate_report_package_count": candidate_report_package_count,
        "ready_report_package_count": ready_report_package_count,
        "sampled_cell_count": metrics["sampled_cell_count"],
        "verified_cell_count": metrics["verified_cell_count"],
        "metrics": metrics,
        "provisional_project_value_result": decision,
        "truth_policy": "only rows with ground_truth_review_status=VERIFIED are counted",
        "readiness_gates": dict(READINESS_GATES_CLOSED),
        "mineru_run_count": 0,
        "ocr_run_count": 0,
        "llm_api_call_count": 0,
        "vlm_api_call_count": 0,
    }


def decide_provisional_project_value(
    metrics: dict[str, Any],
    *,
    ready_report_package_count: int = 1,
) -> str:
    if ready_report_package_count < 1:
        return BENCHMARK_METHOD_INVALID
    if metrics.get("sampled_cell_count", 0) < 1:
        return NEEDS_MORE_VERIFIED_CELLS
    if metrics.get("verified_cell_count", 0) < 20:
        return NEEDS_MORE_VERIFIED_CELLS
    true_positive_count = metrics.get("datefac_true_positive_count", 0)
    false_positive_count = metrics.get("datefac_false_positive_count", 0)
    high_severity_count = metrics.get("high_severity_error_count", 0)
    independent_error_count = metrics.get("mineru_error_count", 0) + metrics.get("original_error_count", 0)
    if independent_error_count == 0 or (true_positive_count == 0 and high_severity_count == 0):
        return ONE_REPORT_SUGGESTS_LOW_VALUE
    if true_positive_count > 0 and metrics.get("precision", 0.0) >= 0.5 and false_positive_count <= true_positive_count:
        return ONE_REPORT_SUGGESTS_LIGHTWEIGHT_QA_VALUE
    return BENCHMARK_METHOD_VALID


def write_summary_files(
    summary: dict[str, Any],
    *,
    json_path: str | Path,
    markdown_path: str | Path,
) -> tuple[Path, Path]:
    json_output_path = Path(json_path)
    markdown_output_path = Path(markdown_path)
    json_output_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_output_path.parent.mkdir(parents=True, exist_ok=True)
    json_output_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=JSON_INDENT) + "\n",
        encoding="utf-8",
    )
    markdown_output_path.write_text(render_summary_markdown(summary), encoding="utf-8")
    return json_output_path, markdown_output_path


def render_summary_markdown(summary: dict[str, Any]) -> str:
    metrics = summary.get("metrics", {})
    lines = [
        "# R7CF Project Necessity Benchmark Summary",
        "",
        f"- task_id: {summary.get('task_id')}",
        f"- benchmark_version: {summary.get('benchmark_version')}",
        f"- candidate_report_package_count: {summary.get('candidate_report_package_count')}",
        f"- ready_report_package_count: {summary.get('ready_report_package_count')}",
        f"- sampled_cell_count: {summary.get('sampled_cell_count')}",
        f"- verified_cell_count: {summary.get('verified_cell_count')}",
        f"- provisional_project_value_result: {summary.get('provisional_project_value_result')}",
        f"- truth_policy: {summary.get('truth_policy')}",
        f"- readiness_gates: {summary.get('readiness_gates')}",
        "",
        "## Metrics",
        "",
    ]
    for key in (
        "mineru_error_count",
        "original_error_count",
        "both_wrong_same_value_count",
        "both_wrong_different_value_count",
        "datefac_true_positive_count",
        "datefac_false_positive_count",
        "datefac_false_negative_count",
        "datefac_true_negative_count",
        "precision",
        "recall",
        "false_positive_rate",
        "high_severity_error_count",
    ):
        lines.append(f"- {key}: {metrics.get(key)}")
    lines.extend(
        [
            "",
            "## Boundary",
            "",
            "- MinerU/OCR/LLM/VLM runs: 0",
            "- clean_data writes: 0",
            "- readiness gates: CLOSED",
            "- generated files are local benchmark artifacts and must not be committed",
            "",
        ]
    )
    return "\n".join(lines)


def resolve_correctness(row: dict[str, Any], *, system: str) -> bool:
    explicit_value = row.get(f"{system}_correct")
    explicit_bool = _parse_bool(explicit_value)
    if explicit_bool is not None:
        return explicit_bool
    system_value = normalize_benchmark_value(row.get(f"{system}_value"))
    truth_value = normalize_benchmark_value(row.get("pdf_ground_truth_value"))
    if not truth_value:
        return False
    return system_value == truth_value and _unit_compatible(row, system=system)


def normalize_benchmark_value(value: Any) -> str:
    normalized = normalize_numeric_value(value)
    if normalized is not None:
        return normalized
    text = "" if value is None else str(value).strip()
    if not text:
        return ""
    try:
        return format(Decimal(text.replace(",", "")).normalize(), "f")
    except InvalidOperation:
        return text


def classify_inventory_file(path: str | Path) -> str | None:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    name = file_path.name.lower()
    if suffix == ".pdf":
        return None if _is_generated_pdf(file_path) else "pdf"
    if suffix == ".json" and "content_list_v2" in name:
        return "mineru"
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return "original"
    if suffix == ".json" and "content_list" not in name and _looks_like_original_json(name):
        return "original"
    return None


def extract_report_id(path: str | Path) -> str | None:
    file_path = Path(path)
    hint = KNOWN_ORIGINAL_ID_HINTS.get(file_path.name.lower())
    if hint:
        return hint
    for part in (file_path.name, *[parent.name for parent in file_path.parents]):
        match = REPORT_ID_RE.search(part)
        if match:
            return match.group(0)
    return None


def is_forbidden_committed_benchmark_artifact(path: str | Path) -> bool:
    file_path = Path(path)
    lower_parts = {part.lower() for part in file_path.parts}
    suffix = file_path.suffix.lower()
    name = file_path.name.lower()
    if "output" in lower_parts or "temp" in lower_parts or "data" in lower_parts:
        return True
    if suffix in FORBIDDEN_COMMITTED_SUFFIXES:
        return True
    return suffix == ".json" and ("content_list" in name or "mineru" in name)


def _iter_inventory_files(root: Path) -> Iterable[Path]:
    ignored_dirs = {".git", "__pycache__", ".pytest_cache", ".mypy_cache", "node_modules"}
    for current_root, dirnames, filenames in os.walk(root):
        dirnames[:] = [dirname for dirname in dirnames if dirname not in ignored_dirs]
        for filename in filenames:
            file_path = Path(current_root) / filename
            if classify_inventory_file(file_path):
                yield file_path


def _dedupe_paths(paths: Sequence[str | Path]) -> tuple[str, ...]:
    deduped: dict[str, str] = {}
    for path in paths:
        file_path = Path(path)
        try:
            key = str(file_path.resolve()).lower()
            value = str(file_path.resolve())
        except OSError:
            key = str(file_path.absolute()).lower()
            value = str(file_path.absolute())
        deduped.setdefault(key, value)
    return tuple(sorted(deduped.values(), key=str.lower))


def _select_preferred_path(paths: Sequence[str], preferred_paths: Sequence[str]) -> str | None:
    if not paths:
        return None
    normalized_paths = {str(Path(path).resolve()).lower(): str(Path(path).resolve()) for path in paths}
    for preferred_path in preferred_paths:
        normalized_preferred = str(Path(preferred_path).resolve()).lower()
        if normalized_preferred in normalized_paths:
            return normalized_paths[normalized_preferred]
    return sorted(paths, key=str.lower)[0]


def _sample_sort_key(row: dict[str, Any]) -> tuple[int, int, int, str, str, str]:
    metric_key = str(row.get("metric_key", row.get("metric", ""))).lower()
    period = str(row.get("period", ""))
    suffix = period[-1:] if period[-1:].isalpha() else ""
    try:
        metric_priority = HIGH_VALUE_METRIC_PRIORITIES.index(metric_key)
    except ValueError:
        metric_priority = len(HIGH_VALUE_METRIC_PRIORITIES)
    review_priority = 0 if row.get("review_required") else 1
    return (
        metric_priority,
        PERIOD_SUFFIX_PRIORITY.get(suffix.upper(), 9),
        review_priority,
        str(row.get("statement_context", "")),
        period,
        str(row.get("status", "")),
    )


def _trace_locator_or_bbox(trace: dict[str, Any]) -> str:
    locator = trace.get("locator")
    bbox = trace.get("bbox")
    if locator and bbox:
        return f"{locator};bbox:{bbox}"
    if locator:
        return str(locator)
    if bbox:
        return f"bbox:{bbox}"
    return ""


def _safe_mapping(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _excel_safe(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _is_verified(value: Any) -> bool:
    return str(value or "").strip().upper() == "VERIFIED"


def _parse_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    text = str(value or "").strip().upper()
    if text in {"TRUE", "YES", "Y", "1"}:
        return True
    if text in {"FALSE", "NO", "N", "0"}:
        return False
    return None


def _unit_compatible(row: dict[str, Any], *, system: str) -> bool:
    truth_unit = str(row.get("pdf_ground_truth_unit", "") or "").strip().lower()
    if not truth_unit:
        return True
    system_unit = str(row.get("unit", "") if system in {"mineru", "original"} else "").strip().lower()
    return not system_unit or system_unit == truth_unit


def _is_high_severity(value: Any) -> bool:
    return str(value or "").strip().upper() in {"HIGH", "CRITICAL", "P0", "P1"}


def _safe_ratio(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return round(numerator / denominator, 6)


def _is_generated_pdf(path: Path) -> bool:
    name = path.name.lower()
    return any(token in name for token in ("_origin.pdf", "_layout.pdf", "_span.pdf"))


def _looks_like_original_json(name: str) -> bool:
    return any(token in name for token in ("original", "datefac", "extract", "extraction", "artifact"))


def stable_hash(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


__all__ = [
    "AMBIGUOUS_PAIRING",
    "ANJING_REPORT_ID",
    "BENCHMARK_METHOD_INVALID",
    "BENCHMARK_METHOD_VALID",
    "BENCHMARK_VERSION",
    "DEFAULT_INVENTORY_ROOTS",
    "MISSING_MINERU_ARTIFACT",
    "MISSING_ORIGINAL_ARTIFACT",
    "MISSING_PDF",
    "NEEDS_MORE_VERIFIED_CELLS",
    "ONE_REPORT_SUGGESTS_LIGHTWEIGHT_QA_VALUE",
    "ONE_REPORT_SUGGESTS_LOW_VALUE",
    "READY",
    "READINESS_GATES_CLOSED",
    "REVIEW_PACK_COLUMNS",
    "ReportPackageCandidate",
    "build_benchmark_summary",
    "build_one_report_review_pack_rows",
    "build_review_pack_row",
    "calculate_benchmark_metrics",
    "classify_inventory_file",
    "classify_report_package",
    "decide_provisional_project_value",
    "extract_report_id",
    "inventory_report_packages",
    "is_forbidden_committed_benchmark_artifact",
    "normalize_benchmark_value",
    "read_review_pack",
    "render_summary_markdown",
    "resolve_correctness",
    "resolve_canonical_report_package",
    "sample_high_value_cells",
    "select_one_report_package",
    "stable_hash",
    "validate_review_pack_schema",
    "write_review_pack",
    "write_summary_files",
]
