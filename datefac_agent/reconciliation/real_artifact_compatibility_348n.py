"""R7CD real MinerU artifact compatibility slice.

This module reads real-style MinerU content_list_v2 table blocks and the
existing DateFac workbook shape into a shared reconciliation record model.
It is intentionally demo-only: no clean_data writes, no DB persistence, no
delivery/export integration, and no readiness changes.
"""

from __future__ import annotations

from copy import deepcopy
from decimal import Decimal, InvalidOperation
import hashlib
from html.parser import HTMLParser
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MATCH = "MATCH"
CONFLICT = "CONFLICT"
MINERU_ONLY = "MINERU_ONLY"
ORIGINAL_ONLY = "ORIGINAL_ONLY"
UNPARSEABLE = "UNPARSEABLE"
UNIT_REVIEW = "UNIT_REVIEW"

REVIEW_REQUIRED_STATUSES = frozenset({CONFLICT, MINERU_ONLY, ORIGINAL_ONLY, UNPARSEABLE, UNIT_REVIEW})
EVIDENCE_PREVIEW_LIMIT = 180
RECONCILIATION_VERSION = "r7cd_real_artifact_compatibility_slice_v1"

READINESS_GATES_CLOSED: dict[str, bool] = {
    "client_ready": False,
    "production_ready": False,
    "formal_client_export_allowed": False,
    "demo_export_only": True,
}

TARGET_EXCEL_SHEETS: tuple[str, ...] = (
    "Financial_Data_Valuation",
    "Balance_Sheet",
    "Income_Statement",
    "Cash_Flow",
    "Ratios_Per_Share",
)

EXCEL_SHEET_CONTEXT: dict[str, str] = {
    "Financial_Data_Valuation": "financial_data_valuation",
    "Balance_Sheet": "balance_sheet",
    "Income_Statement": "income_statement",
    "Cash_Flow": "cash_flow",
    "Ratios_Per_Share": "ratios_per_share",
}

CONTEXT_DISPLAY_NAMES: dict[str, str] = {
    "financial_data_valuation": "Financial Data & Valuation",
    "balance_sheet": "Balance Sheet",
    "income_statement": "Income Statement",
    "cash_flow": "Cash Flow",
    "ratios_per_share": "Ratios & Per Share",
    "unknown": "Unknown",
}

UNIT_ALIASES: dict[str, str] = {
    "百万元": "cny_million",
    "人民币百万元": "cny_million",
    "cny_m": "cny_million",
    "cnym": "cny_million",
    "亿元": "cny_100_million",
    "亿股": "shares_100_million",
    "%": "percent",
    "pct": "percent",
    "百分比": "percent",
    "倍": "multiple",
    "次": "times",
    "元": "yuan",
    "元/股": "yuan_per_share",
    "元每股": "yuan_per_share",
}

METRIC_ALIAS_BY_COMPACT_LABEL: dict[str, str] = {
    "营业收入": "revenue",
    "收入": "revenue",
    "营业成本": "operating_cost",
    "营业税金及附加": "taxes_and_surcharges",
    "营业费用": "selling_expenses",
    "管理费用": "administrative_expenses",
    "研发费用": "r_and_d_expenses",
    "财务费用": "financial_expenses",
    "资产减值损失": "asset_impairment_loss",
    "公允价值变动收益": "fair_value_change_income",
    "投资净收益": "investment_income",
    "投资损失": "investment_loss",
    "营业利润": "operating_profit",
    "营业外收入": "non_operating_income",
    "营业外支出": "non_operating_expenses",
    "利润总额": "total_profit",
    "所得税": "income_tax",
    "税后利润": "profit_after_tax",
    "少数股东损益": "minority_interest",
    "归属母公司净利润": "parent_net_profit",
    "归属于母公司净利润": "parent_net_profit",
    "净利润": "net_profit",
    "ebitda": "ebitda",
    "eps摊薄": "eps_diluted",
    "eps": "eps",
    "每股收益最新摊薄": "eps_latest_diluted",
    "每股经营现金流最新摊薄": "operating_cash_flow_per_share_latest_diluted",
    "每股净资产最新摊薄": "book_value_per_share_latest_diluted",
    "roe": "roe",
    "roic": "roic",
    "pe": "pe",
    "p/e": "pe",
    "pb": "pb",
    "p/b": "pb",
    "ev/ebitda": "ev_ebitda",
    "毛利率": "gross_margin",
    "净利率": "net_margin",
    "营业收入增长率": "revenue_growth",
    "营业利润增长率": "operating_profit_growth",
    "归母净利润增长率": "parent_net_profit_growth",
    "归属母公司净利润增长率": "parent_net_profit_growth",
    "归属于母公司净利润增长率": "parent_net_profit_growth",
    "资产负债率": "debt_to_asset_ratio",
    "流动比率": "current_ratio",
    "速动比率": "quick_ratio",
    "总资产周转率": "total_asset_turnover",
    "应收账款周转率": "accounts_receivable_turnover",
    "应付账款周转率": "accounts_payable_turnover",
    "经营活动现金流": "operating_cash_flow",
    "投资活动现金流": "investing_cash_flow",
    "筹资活动现金流": "financing_cash_flow",
}

PERIOD_RE = re.compile(r"(?P<year>(?:19|20)\d{2})(?P<suffix>[AEF])?", re.IGNORECASE)


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell_parts: list[str] | None = None
        self._cell_colspan = 1

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "tr":
            self._row = []
        elif tag in {"td", "th"}:
            self._cell_parts = []
            self._cell_colspan = _safe_int(dict(attrs).get("colspan"), default=1)
        elif tag == "br" and self._cell_parts is not None:
            self._cell_parts.append(" ")

    def handle_data(self, data: str) -> None:
        if self._cell_parts is not None:
            self._cell_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"td", "th"} and self._row is not None and self._cell_parts is not None:
            cell = _compact_text("".join(self._cell_parts))
            for _ in range(max(self._cell_colspan, 1)):
                self._row.append(cell)
            self._cell_parts = None
            self._cell_colspan = 1
        elif tag == "tr" and self._row is not None:
            if any(cell for cell in self._row):
                self.rows.append(self._row)
            self._row = None


def load_json_artifact(path: str | Path) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def expand_html_table(html: Any) -> list[list[str]]:
    if not isinstance(html, str) or not html.strip():
        return []
    parser = _HTMLTableParser()
    parser.feed(html)
    return parser.rows


def extract_mineru_records(mineru_artifact: Any) -> list[dict[str, Any]]:
    artifact = deepcopy(mineru_artifact)
    pages = _extract_mineru_pages(artifact)
    records: list[dict[str, Any]] = []
    for page_idx, page_blocks in enumerate(pages):
        if not isinstance(page_blocks, list):
            continue
        page_number = page_idx + 1
        for block_index, block in enumerate(page_blocks):
            if not isinstance(block, dict) or str(block.get("type", "")).lower() != "table":
                continue
            content = block.get("content") if isinstance(block.get("content"), dict) else {}
            html = content.get("html") if isinstance(content, dict) else None
            matrix = expand_html_table(html)
            if not matrix:
                continue
            caption_text = _flatten_content_text(content.get("table_caption") if isinstance(content, dict) else None)
            footnote_text = _flatten_content_text(content.get("table_footnote") if isinstance(content, dict) else None)
            table_context = determine_statement_context(caption_text)
            default_unit = _unit_from_text(caption_text)
            rows = _records_from_matrix(
                source="mineru",
                matrix=matrix,
                statement_context=table_context,
                default_unit=default_unit,
                source_trace_base={
                    "source": "mineru",
                    "page_number": page_number,
                    "page_idx": page_idx,
                    "block_index": block_index,
                    "bbox": deepcopy(block.get("bbox")),
                    "table_caption": caption_text,
                    "table_footnote": footnote_text,
                },
            )
            records.extend(rows)
    return records


def load_original_xlsx_records(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return extract_original_records_from_workbook(workbook)
    finally:
        workbook.close()


def extract_original_records_from_workbook(workbook: Any) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for sheet_name in TARGET_EXCEL_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
        statement_context = EXCEL_SHEET_CONTEXT[sheet_name]
        records.extend(
            _records_from_matrix(
                source="original",
                matrix=matrix,
                statement_context=statement_context,
                default_unit=None,
                source_trace_base={
                    "source": "original",
                    "sheet": sheet_name,
                },
            )
        )
    return records


def compare_records(
    mineru_records: list[dict[str, Any]],
    original_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    mineru_by_key, mineru_unparseable = _index_records(mineru_records, source_label="MinerU")
    original_by_key, original_unparseable = _index_records(original_records, source_label="Original")
    comparison_rows: list[dict[str, Any]] = []

    for key in sorted(set(mineru_by_key) | set(original_by_key), key=_sort_record_key):
        mineru_record = mineru_by_key.get(key)
        original_record = original_by_key.get(key)
        if mineru_record and original_record:
            if _units_incompatible(mineru_record, original_record):
                comparison_rows.append(
                    _comparison_row(
                        UNIT_REVIEW,
                        "unit metadata differs for the same statement, metric, and period",
                        mineru_record,
                        original_record,
                    )
                )
            elif mineru_record["normalized_value"] == original_record["normalized_value"]:
                comparison_rows.append(
                    _comparison_row(
                        MATCH,
                        "values match after statement, metric, period, numeric, and unit normalization",
                        mineru_record,
                        original_record,
                    )
                )
            else:
                comparison_rows.append(
                    _comparison_row(
                        CONFLICT,
                        "normalized values differ for the same statement, metric, and period",
                        mineru_record,
                        original_record,
                    )
                )
        elif mineru_record:
            comparison_rows.append(
                _comparison_row(
                    MINERU_ONLY,
                    "statement-metric-period exists only in MinerU artifact",
                    mineru_record,
                    None,
                )
            )
        elif original_record:
            comparison_rows.append(
                _comparison_row(
                    ORIGINAL_ONLY,
                    "statement-metric-period exists only in original workbook",
                    None,
                    original_record,
                )
            )

    for record in mineru_unparseable:
        comparison_rows.append(_comparison_row(UNPARSEABLE, record["unparseable_reason"], record, None))
    for record in original_unparseable:
        comparison_rows.append(_comparison_row(UNPARSEABLE, record["unparseable_reason"], None, record))

    return sorted(comparison_rows, key=_comparison_sort_key)


def reconcile_real_artifacts(mineru_artifact: Any, original_workbook: Any) -> dict[str, Any]:
    mineru_records = extract_mineru_records(mineru_artifact)
    original_records = extract_original_records_from_workbook(original_workbook)
    comparison_rows = compare_records(mineru_records, original_records)
    review_candidates = build_review_queue_candidates(comparison_rows)
    return _result(mineru_records, original_records, comparison_rows, review_candidates)


def load_and_reconcile_real_artifacts(mineru_json_path: str | Path, original_xlsx_path: str | Path) -> dict[str, Any]:
    mineru_artifact = load_json_artifact(mineru_json_path)
    workbook = load_workbook(original_xlsx_path, read_only=True, data_only=True)
    try:
        return reconcile_real_artifacts(mineru_artifact, workbook)
    finally:
        workbook.close()


def build_review_queue_candidates(comparison_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for row in comparison_rows:
        if not row["review_required"]:
            continue
        identity = "|".join(
            [
                RECONCILIATION_VERSION,
                row["statement_context"],
                row["metric_key"],
                row["period"],
                row["status"],
            ]
        )
        candidates.append(
            {
                "review_item_id": "r7cd:" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24],
                "statement_context": row["statement_context"],
                "candidate_metric_name": row["metric"],
                "candidate_period": row["period"],
                "agreement_status": row["status"],
                "unit_review_required": row["status"] == UNIT_REVIEW,
                "reason": row["reason"],
                "blocked_delivery_reason": row["blocked_delivery_reason"],
                "evidence_preview": row["evidence_preview"],
                "source_trace": deepcopy(row["source_trace"]),
                "review_status": "PENDING_REVIEW",
                "clean_data_eligible": False,
                "readiness_gates": deepcopy(READINESS_GATES_CLOSED),
            }
        )
    return candidates


def determine_statement_context(text: Any) -> str:
    compact = _normalize_label_text(_flatten_content_text(text))
    if "资产负债表" in compact:
        return "balance_sheet"
    if "利润表" in compact:
        return "income_statement"
    if "现金流量表" in compact or "现金流量" in compact:
        return "cash_flow"
    if "主要财务比率" in compact or "每股指标" in compact:
        return "ratios_per_share"
    if "财务数据与估值" in compact or "估值" in compact:
        return "financial_data_valuation"
    return "unknown"


def normalize_period(period: Any) -> str | None:
    text = _compact_text(period)
    if not text:
        return None
    match = PERIOD_RE.search(text)
    if not match:
        return text
    suffix = (match.group("suffix") or "").upper()
    return match.group("year") + suffix


def normalize_numeric_value(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value)
    text = _compact_text(value)
    if not text:
        return None
    text = text.replace(",", "").replace("，", "").replace(" ", "").replace("%", "")
    if re.fullmatch(r"[+-]?\d+\.\d+\.\d+", text):
        parts = text.split(".")
        if len(parts) == 3 and len(parts[-1]) == 1:
            text = parts[0] + "." + parts[-1]
    is_parenthesized_negative = text.startswith("(") and text.endswith(")")
    if is_parenthesized_negative:
        text = "-" + text[1:-1]
    text = re.sub(r"(百万元|亿元|万元|元/股|元|倍|次)$", "", text, flags=re.IGNORECASE)
    try:
        decimal = Decimal(text)
    except InvalidOperation:
        return None
    if decimal == 0:
        return "0"
    return format(decimal.normalize(), "f")


def normalize_metric_and_unit(
    raw_metric: Any,
    raw_unit: Any = None,
    *,
    statement_context: str = "unknown",
    section_label: str | None = None,
    previous_metric_key: str | None = None,
    previous_metric_label: str | None = None,
) -> tuple[str, str, str | None, str | None]:
    metric_label, embedded_unit = _split_metric_label(raw_metric)
    raw_unit_text = _compact_text(raw_unit) or embedded_unit
    if not raw_unit_text and section_label:
        raw_unit_text = _unit_from_text(section_label)

    compact_label = _normalize_label_text(metric_label)
    if compact_label.lower() == "yoy" and previous_metric_key and previous_metric_label:
        metric_key = previous_metric_key + "_yoy"
        display_label = previous_metric_label + " YoY"
        raw_unit_text = raw_unit_text or "%"
    else:
        metric_key = _metric_key_from_label(metric_label, raw_unit_text, statement_context)
        display_label = metric_label

    raw_unit_text = raw_unit_text or _infer_unit(metric_key, display_label, statement_context)
    normalized_unit = normalize_unit(raw_unit_text, metric_key=metric_key)
    return metric_key, display_label, raw_unit_text or None, normalized_unit


def normalize_unit(unit: Any, *, metric_key: str | None = None) -> str | None:
    text = _normalize_unit_text(unit)
    if not text:
        return None
    if text.startswith("亿") and text != "亿元":
        return None
    if metric_key and ("per_share" in metric_key or metric_key.startswith("eps")) and "元" in text:
        return "yuan_per_share"
    if text in UNIT_ALIASES:
        return UNIT_ALIASES[text]
    if "百万元" in text:
        return "cny_million"
    if "亿元" in text:
        return "cny_100_million"
    if "%" in text or "pct" in text.lower():
        return "percent"
    if "倍" in text:
        return "multiple"
    if "次" in text:
        return "times"
    if "元" in text:
        return "yuan"
    return text.lower()


def _records_from_matrix(
    *,
    source: str,
    matrix: list[list[Any]],
    statement_context: str,
    default_unit: str | None,
    source_trace_base: dict[str, Any],
) -> list[dict[str, Any]]:
    header_index, period_columns, unit_column_index = _find_header_row(matrix)
    if header_index is None:
        return []

    records: list[dict[str, Any]] = []
    current_section: str | None = None
    current_section_unit: str | None = None
    previous_metric_key: str | None = None
    previous_metric_label: str | None = None

    for matrix_row_index in range(header_index + 1, len(matrix)):
        row = matrix[matrix_row_index]
        raw_metric = row[0] if row else None
        metric_text = _compact_text(raw_metric)
        if not metric_text:
            continue
        period_values = [_cell_at(row, column_index) for column_index, _period in period_columns]
        if not any(_compact_text(value) for value in period_values):
            current_section = metric_text
            current_section_unit = _unit_from_text(metric_text)
            continue

        row_context = _context_for_row(statement_context, current_section, metric_text)
        row_unit = _cell_at(row, unit_column_index) if unit_column_index is not None else None
        row_unit = _compact_text(row_unit) or default_unit or current_section_unit
        metric_key, metric_label, raw_unit, normalized_unit = normalize_metric_and_unit(
            raw_metric,
            row_unit,
            statement_context=row_context,
            section_label=current_section,
            previous_metric_key=previous_metric_key,
            previous_metric_label=previous_metric_label,
        )

        for column_index, period in period_columns:
            raw_value = _cell_at(row, column_index)
            normalized_value = normalize_numeric_value(raw_value)
            record = {
                "source": source,
                "statement_context": row_context,
                "statement_context_display": CONTEXT_DISPLAY_NAMES.get(row_context, row_context),
                "metric_key": metric_key,
                "metric": metric_label,
                "period": period,
                "normalized_value": normalized_value,
                "raw_metric": raw_metric,
                "raw_period": period,
                "raw_value": raw_value,
                "unit": raw_unit,
                "normalized_unit": normalized_unit,
                "evidence_preview": _record_preview(row_context, metric_label, period, raw_value, raw_unit),
                "source_trace": _source_trace(source_trace_base, matrix_row_index, column_index, period),
            }
            records.append(record)

        if not _is_yoy_label(metric_text):
            previous_metric_key = metric_key
            previous_metric_label = metric_label

    return records


def _find_header_row(matrix: list[list[Any]]) -> tuple[int | None, list[tuple[int, str]], int | None]:
    for row_index, row in enumerate(matrix):
        period_columns = []
        unit_column_index = None
        for column_index, cell in enumerate(row):
            period = normalize_period(cell)
            if period and PERIOD_RE.fullmatch(_compact_text(cell)):
                period_columns.append((column_index, period))
            header_text = _normalize_label_text(cell).replace("/", "").replace("／", "")
            if header_text in {"单位口径", "单位"}:
                unit_column_index = column_index
        if len(period_columns) >= 2:
            return row_index, period_columns, unit_column_index
    return None, [], None


def _source_trace(
    source_trace_base: dict[str, Any],
    matrix_row_index: int,
    column_index: int,
    period: str,
) -> dict[str, Any]:
    trace = deepcopy(source_trace_base)
    trace["row_index"] = matrix_row_index
    trace["row_number"] = matrix_row_index + 1
    trace["column_index"] = column_index
    trace["period"] = period
    if trace.get("source") == "mineru":
        trace["locator"] = f"page:{trace.get('page_number')}:block:{trace.get('block_index')}:row:{matrix_row_index}:col:{column_index}"
    elif trace.get("source") == "original":
        trace["locator"] = f"sheet:{trace.get('sheet')}:row:{matrix_row_index + 1}:col:{column_index + 1}"
    return {key: value for key, value in trace.items() if value not in (None, "", [])}


def _index_records(
    records: list[dict[str, Any]],
    *,
    source_label: str,
) -> tuple[dict[tuple[str, str, str], dict[str, Any]], list[dict[str, Any]]]:
    indexed: dict[tuple[str, str, str], dict[str, Any]] = {}
    unparseable: list[dict[str, Any]] = []
    for record in records:
        if not record.get("metric_key") or not record.get("period") or record.get("normalized_value") is None:
            problem = deepcopy(record)
            problem["unparseable_reason"] = f"{source_label} row could not be normalized"
            unparseable.append(problem)
            continue
        key = _record_key(record)
        if key in indexed:
            duplicate = deepcopy(record)
            duplicate["unparseable_reason"] = f"{source_label} duplicate statement-metric-period key requires review"
            unparseable.append(duplicate)
            continue
        indexed[key] = record
    return indexed, unparseable


def _comparison_row(
    status: str,
    reason: str,
    mineru_record: dict[str, Any] | None,
    original_record: dict[str, Any] | None,
) -> dict[str, Any]:
    primary = mineru_record or original_record or {}
    review_required = status in REVIEW_REQUIRED_STATUSES
    return {
        "statement_context": primary.get("statement_context", "unknown"),
        "statement_context_display": primary.get("statement_context_display", "Unknown"),
        "metric_key": primary.get("metric_key", "unparseable_metric"),
        "metric": primary.get("metric", "UNPARSEABLE_METRIC"),
        "period": primary.get("period") or "UNPARSEABLE_PERIOD",
        "mineru_value": None if mineru_record is None else mineru_record.get("normalized_value"),
        "original_value": None if original_record is None else original_record.get("normalized_value"),
        "mineru_unit": None if mineru_record is None else mineru_record.get("normalized_unit"),
        "original_unit": None if original_record is None else original_record.get("normalized_unit"),
        "status": status,
        "reason": reason,
        "review_required": review_required,
        "blocked_delivery_reason": "" if not review_required else _blocked_delivery_reason(status),
        "evidence_preview": _comparison_preview(mineru_record, original_record),
        "source_trace": {
            "mineru": None if mineru_record is None else deepcopy(mineru_record["source_trace"]),
            "original": None if original_record is None else deepcopy(original_record["source_trace"]),
        },
    }


def _result(
    mineru_records: list[dict[str, Any]],
    original_records: list[dict[str, Any]],
    comparison_rows: list[dict[str, Any]],
    review_candidates: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "reconciliation_version": RECONCILIATION_VERSION,
        "mineru_record_count": len(mineru_records),
        "original_record_count": len(original_records),
        "comparison_rows": comparison_rows,
        "review_queue_candidates": review_candidates,
        "summary": {
            "mineru_record_count": len(mineru_records),
            "original_record_count": len(original_records),
            "comparison_row_count": len(comparison_rows),
            "match_count": sum(1 for row in comparison_rows if row["status"] == MATCH),
            "conflict_count": sum(1 for row in comparison_rows if row["status"] == CONFLICT),
            "mineru_only_count": sum(1 for row in comparison_rows if row["status"] == MINERU_ONLY),
            "original_only_count": sum(1 for row in comparison_rows if row["status"] == ORIGINAL_ONLY),
            "unparseable_count": sum(1 for row in comparison_rows if row["status"] == UNPARSEABLE),
            "unit_review_count": sum(1 for row in comparison_rows if row["status"] == UNIT_REVIEW),
            "review_required_count": len(review_candidates),
            "readiness_gates": deepcopy(READINESS_GATES_CLOSED),
        },
    }


def _record_key(record: dict[str, Any]) -> tuple[str, str, str]:
    return (record["statement_context"], record["metric_key"], record["period"])


def _units_incompatible(mineru_record: dict[str, Any], original_record: dict[str, Any]) -> bool:
    mineru_unit = mineru_record.get("normalized_unit")
    original_unit = original_record.get("normalized_unit")
    return bool(mineru_unit and original_unit and mineru_unit != original_unit)


def _blocked_delivery_reason(status: str) -> str:
    return {
        CONFLICT: "value_conflict_requires_review",
        MINERU_ONLY: "mineru_only_requires_review",
        ORIGINAL_ONLY: "original_only_requires_review",
        UNPARSEABLE: "unparseable_row_requires_review",
        UNIT_REVIEW: "unit_mismatch_requires_review",
    }.get(status, "")


def _comparison_preview(
    mineru_record: dict[str, Any] | None,
    original_record: dict[str, Any] | None,
) -> str:
    previews: list[str] = []
    if mineru_record is not None:
        previews.append("MinerU: " + mineru_record["evidence_preview"])
    if original_record is not None:
        previews.append("Original: " + original_record["evidence_preview"])
    return _bounded_preview(" | ".join(previews))


def _record_preview(context: str, metric: str, period: str, value: Any, unit: Any) -> str:
    unit_text = _compact_text(unit)
    unit_suffix = f" {unit_text}" if unit_text else ""
    return _bounded_preview(f"{CONTEXT_DISPLAY_NAMES.get(context, context)} | {metric} | {period} | {value}{unit_suffix}")


def _bounded_preview(value: Any, *, limit: int = EVIDENCE_PREVIEW_LIMIT) -> str:
    text = _compact_text(value)
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def _extract_mineru_pages(artifact: Any) -> list[Any]:
    if isinstance(artifact, list):
        return artifact
    if isinstance(artifact, dict):
        pages = artifact.get("content_list_v2", artifact.get("pages"))
        if isinstance(pages, list):
            return pages
    return []


def _context_for_row(table_context: str, section_label: str | None, metric_text: str) -> str:
    section_compact = _normalize_label_text(section_label)
    metric_compact = _normalize_label_text(metric_text)
    if table_context == "cash_flow" and ("每股指标" in section_compact or "每股" in metric_compact):
        return "ratios_per_share"
    return table_context


def _metric_key_from_label(metric_label: str, raw_unit: str | None, statement_context: str) -> str:
    compact = _normalize_label_text(metric_label)
    compact_lower = compact.lower()
    normalized_unit = normalize_unit(raw_unit)
    if statement_context == "ratios_per_share" and normalized_unit == "percent":
        if compact in {"营业收入", "收入"}:
            return "revenue_growth"
        if compact == "营业利润":
            return "operating_profit_growth"
        if compact in {"归属于母公司净利润", "归属母公司净利润", "归母净利润"}:
            return "parent_net_profit_growth"
    if compact in METRIC_ALIAS_BY_COMPACT_LABEL:
        return METRIC_ALIAS_BY_COMPACT_LABEL[compact]
    if compact_lower in METRIC_ALIAS_BY_COMPACT_LABEL:
        return METRIC_ALIAS_BY_COMPACT_LABEL[compact_lower]
    return compact_lower or "unparseable_metric"


def _split_metric_label(raw_metric: Any) -> tuple[str, str | None]:
    text = _compact_text(raw_metric).replace("（", "(").replace("）", ")")
    if not text:
        return "", None
    match = re.fullmatch(r"(?P<base>.*?)\((?P<inside>[^()]*)\)\s*[:：]?", text)
    if not match:
        return text.rstrip(":："), None
    base = _compact_text(match.group("base")).rstrip(":：")
    inside = _compact_text(match.group("inside"))
    unit = _unit_from_text(inside)
    qualifier = _strip_unit_words(inside)
    if qualifier:
        base = _compact_text(f"{base} {qualifier}")
    return base, unit


def _strip_unit_words(text: str) -> str:
    value = text
    for token in ("百万元", "亿元", "万元", "元/股", "元", "%", "倍", "次"):
        value = value.replace(token, "")
    value = re.sub(r"[/,，;；、]+", " ", value)
    return _compact_text(value)


def _unit_from_text(text: Any) -> str | None:
    compact = _compact_text(text)
    if not compact:
        return None
    for token in ("百万元", "亿元", "万元", "元/股", "元", "%", "倍", "次"):
        if token in compact:
            return token
    return None


def _infer_unit(metric_key: str, metric_label: str, statement_context: str) -> str | None:
    compact = _normalize_label_text(metric_label)
    if metric_key.endswith("_growth") or metric_key.endswith("_yoy"):
        return "%"
    if metric_key in {"gross_margin", "net_margin", "roe", "roic", "debt_to_asset_ratio"}:
        return "%"
    if metric_key in {"pe", "pb"} or "比率" in compact:
        return "倍"
    if "周转率" in compact:
        return "次"
    if "per_share" in metric_key or metric_key.startswith("eps"):
        return "元"
    if statement_context in {"balance_sheet", "income_statement", "cash_flow"}:
        return "百万元"
    return None


def _normalize_unit_text(value: Any) -> str:
    text = _compact_text(value).replace("（", "(").replace("）", ")")
    text = text.replace("／", "/")
    return text.rstrip(":：")


def _normalize_label_text(value: Any) -> str:
    text = _compact_text(value)
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"[\s:：,，;；、()（）]", "", text)
    return text


def _is_yoy_label(value: Any) -> bool:
    return _normalize_label_text(value).lower() == "yoy"


def _flatten_content_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return _compact_text(value)
    if isinstance(value, dict):
        if "content" in value:
            return _flatten_content_text(value["content"])
        return _compact_text(" ".join(_flatten_content_text(item) for item in value.values()))
    if isinstance(value, list):
        return _compact_text(" ".join(_flatten_content_text(item) for item in value))
    return _compact_text(value)


def _compact_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _cell_at(row: list[Any], column_index: int | None) -> Any:
    if column_index is None or column_index >= len(row):
        return None
    return row[column_index]


def _safe_int(value: Any, *, default: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _sort_record_key(key: tuple[str, str, str]) -> tuple[str, str, str]:
    return key


def _comparison_sort_key(row: dict[str, Any]) -> tuple[str, str, str, str]:
    return (row["statement_context"], row["metric_key"], row["period"], row["status"])


__all__ = [
    "CONFLICT",
    "EVIDENCE_PREVIEW_LIMIT",
    "MATCH",
    "MINERU_ONLY",
    "ORIGINAL_ONLY",
    "READINESS_GATES_CLOSED",
    "RECONCILIATION_VERSION",
    "REVIEW_REQUIRED_STATUSES",
    "TARGET_EXCEL_SHEETS",
    "UNPARSEABLE",
    "UNIT_REVIEW",
    "build_review_queue_candidates",
    "compare_records",
    "determine_statement_context",
    "expand_html_table",
    "extract_mineru_records",
    "extract_original_records_from_workbook",
    "load_and_reconcile_real_artifacts",
    "load_json_artifact",
    "load_original_xlsx_records",
    "normalize_metric_and_unit",
    "normalize_numeric_value",
    "normalize_period",
    "normalize_unit",
    "reconcile_real_artifacts",
]
