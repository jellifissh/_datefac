from __future__ import annotations

import ast
from copy import deepcopy
import json
from pathlib import Path

from openpyxl import Workbook

from datefac_agent.reconciliation.real_artifact_compatibility_348n import (
    CONFLICT,
    EVIDENCE_PREVIEW_LIMIT,
    MATCH,
    MINERU_ONLY,
    ORIGINAL_ONLY,
    READINESS_GATES_CLOSED,
    UNIT_REVIEW,
    build_review_queue_candidates,
    compare_records,
    expand_html_table,
    extract_mineru_records,
    extract_original_records_from_workbook,
    load_and_reconcile_real_artifacts,
    normalize_metric_and_unit,
    normalize_numeric_value,
    normalize_period,
    reconcile_real_artifacts,
)

FIXTURE_PATH = Path("tests/agent/fixtures/mineru_original_reconciliation/real_content_list_v2_table_sample.json")
MODULE_PATH = Path("datefac_agent/reconciliation/real_artifact_compatibility_348n.py")


def _mineru_fixture() -> list:
    return json.loads(FIXTURE_PATH.read_text(encoding="utf-8"))


def _workbook(*, conflict: bool = False, incompatible_unit: bool = False) -> Workbook:
    workbook = Workbook()
    financial = workbook.active
    financial.title = "Financial_Data_Valuation"
    financial.append(["财务数据与估值（P2）"])
    financial.append([])
    financial.append(["指标", "2024A", "2025A", "2026E", "单位/口径"])
    financial.append(["营业收入", 15127, 16193, 18379, "亿元" if incompatible_unit else "百万元"])
    financial.append(["YoY", 7.7, 7.0, 13.5, "%"])
    financial.append(["净利润", 1485, 1359, 1792 if conflict else 1791, "百万元"])
    financial.append(["EPS（摊薄）", 4.46, 4.08, 5.37, "元"])
    financial.append(["ROE", 11.4, 8.6, 10.3, "%"])
    financial.append(["P/E", 22.0, 24.1, 18.3, "倍"])

    income = workbook.create_sheet("Income_Statement")
    income.append(["利润表预测（百万元，P3）"])
    income.append([])
    income.append(["指标", "2024A", "2025A", "2026E", "单位/口径"])
    income.append(["营业收入", 15127, 16193, 18379, "百万元"])
    income.append(["归属母公司净利润", 1485, 1359, 1791, "百万元"])

    ratios = workbook.create_sheet("Ratios_Per_Share")
    ratios.append(["主要财务比率与每股指标（P3）"])
    ratios.append([])
    ratios.append(["指标", "2024A", "2025A", "2026E", "单位/口径"])
    ratios.append(["营业收入增长率", 7.7, 7.0, 13.5, "%"])
    ratios.append(["归母净利润增长率", 0.5, -8.5, 31.7, "%"])
    ratios.append(["P/B", 2.5, 2.1, 1.9, "倍"])
    return workbook


def _rows_by_key(result: dict) -> dict[tuple[str, str, str], dict]:
    return {
        (row["statement_context"], row["metric_key"], row["period"]): row
        for row in result["comparison_rows"]
    }


def test_r7cd_fixture_is_small_real_page_grouped_content_list_v2_shape() -> None:
    artifact = _mineru_fixture()

    assert isinstance(artifact, list)
    assert isinstance(artifact[1], list)
    assert FIXTURE_PATH.stat().st_size < 10000
    assert artifact[1][0]["type"] == "table"
    assert "content" in artifact[1][0]
    assert "html" in artifact[1][0]["content"]


def test_r7cd_expands_html_table_matrix() -> None:
    matrix = expand_html_table("<table><tr><th>会计年度</th><th>2026E</th></tr><tr><td>营业收入(百万元)</td><td>18,379</td></tr></table>")

    assert matrix == [["会计年度", "2026E"], ["营业收入(百万元)", "18,379"]]


def test_r7cd_extracts_mineru_records_with_period_context_and_source_trace() -> None:
    records = extract_mineru_records(_mineru_fixture())
    revenue_2026 = next(
        record
        for record in records
        if record["statement_context"] == "financial_data_valuation"
        and record["metric_key"] == "revenue"
        and record["period"] == "2026E"
    )

    assert revenue_2026["normalized_value"] == "18379"
    assert revenue_2026["unit"] == "百万元"
    assert revenue_2026["normalized_unit"] == "cny_million"
    assert revenue_2026["source_trace"]["page_number"] == 2
    assert revenue_2026["source_trace"]["block_index"] == 0
    assert revenue_2026["source_trace"]["bbox"] == [359, 131, 944, 316]
    assert revenue_2026["source_trace"]["locator"] == "page:2:block:0:row:1:col:3"
    assert len(revenue_2026["evidence_preview"]) <= EVIDENCE_PREVIEW_LIMIT


def test_r7cd_extracts_excel_records_from_selected_sheets() -> None:
    records = extract_original_records_from_workbook(_workbook())
    financial_revenue = next(
        record
        for record in records
        if record["statement_context"] == "financial_data_valuation"
        and record["metric_key"] == "revenue"
        and record["period"] == "2026E"
    )
    income_revenue = next(
        record
        for record in records
        if record["statement_context"] == "income_statement"
        and record["metric_key"] == "revenue"
        and record["period"] == "2026E"
    )

    assert financial_revenue["normalized_value"] == "18379"
    assert financial_revenue["source_trace"]["locator"] == "sheet:Financial_Data_Valuation:row:4:col:4"
    assert income_revenue["normalized_value"] == "18379"
    assert financial_revenue["statement_context"] != income_revenue["statement_context"]


def test_r7cd_metric_unit_numeric_and_period_normalization() -> None:
    metric_key, metric_label, unit, normalized_unit = normalize_metric_and_unit(
        "EPS(摊薄/元)",
        statement_context="financial_data_valuation",
    )

    assert normalize_period("2026E") == "2026E"
    assert normalize_period("FY2026E") == "2026E"
    assert normalize_numeric_value("18,379.00") == "18379"
    assert normalize_numeric_value("(12.50%)") == "-12.5"
    assert normalize_numeric_value("0.08.8") == "0.8"
    assert (metric_key, metric_label, unit, normalized_unit) == (
        "eps_diluted",
        "EPS 摊薄",
        "元",
        "yuan_per_share",
    )


def test_r7cd_reconciles_matching_real_style_rows_and_keeps_context_separate() -> None:
    result = reconcile_real_artifacts(_mineru_fixture(), _workbook())
    rows = _rows_by_key(result)

    assert rows[("financial_data_valuation", "revenue", "2026E")]["status"] == MATCH
    assert rows[("income_statement", "revenue", "2026E")]["status"] == MATCH
    assert rows[("ratios_per_share", "revenue_growth", "2026E")]["status"] == MATCH
    assert result["summary"]["match_count"] >= 15
    assert result["summary"]["readiness_gates"] == READINESS_GATES_CLOSED


def test_r7cd_conflict_missing_and_unit_review_statuses_build_review_candidates() -> None:
    result = reconcile_real_artifacts(_mineru_fixture(), _workbook(conflict=True, incompatible_unit=True))
    rows = _rows_by_key(result)

    assert rows[("financial_data_valuation", "net_profit", "2026E")]["status"] == CONFLICT
    assert rows[("financial_data_valuation", "revenue", "2026E")]["status"] == UNIT_REVIEW

    mineru_records = [
        {
            "source": "mineru",
            "statement_context": "financial_data_valuation",
            "statement_context_display": "Financial Data & Valuation",
            "metric_key": "pe",
            "metric": "P/E",
            "period": "2028E",
            "normalized_value": "13.6",
            "normalized_unit": "multiple",
            "evidence_preview": "MinerU P/E 2028E 13.6",
            "source_trace": {"source": "mineru"},
        }
    ]
    original_records = [
        {
            "source": "original",
            "statement_context": "financial_data_valuation",
            "statement_context_display": "Financial Data & Valuation",
            "metric_key": "pb",
            "metric": "P/B",
            "period": "2028E",
            "normalized_value": "1.5",
            "normalized_unit": "multiple",
            "evidence_preview": "Original P/B 2028E 1.5",
            "source_trace": {"source": "original"},
        }
    ]
    comparison_rows = compare_records(mineru_records, original_records)
    statuses = {row["status"] for row in comparison_rows}
    candidates = build_review_queue_candidates(comparison_rows)

    assert {MINERU_ONLY, ORIGINAL_ONLY}.issubset(statuses)
    assert all(candidate["clean_data_eligible"] is False for candidate in result["review_queue_candidates"])
    assert all(candidate["readiness_gates"] == READINESS_GATES_CLOSED for candidate in result["review_queue_candidates"])
    assert len(candidates) == 2


def test_r7cd_bounded_preview_input_mutation_safety_and_determinism() -> None:
    artifact = _mineru_fixture()
    artifact[1][0]["content"]["table_caption"][0]["content"] = "财务数据与估值：" + "x" * 1000
    before = deepcopy(artifact)

    first = reconcile_real_artifacts(artifact, _workbook(conflict=True))
    second = reconcile_real_artifacts(artifact, _workbook(conflict=True))

    assert artifact == before
    assert first == second
    assert all(len(row["evidence_preview"]) <= EVIDENCE_PREVIEW_LIMIT for row in first["comparison_rows"])
    assert all(len(candidate["evidence_preview"]) <= EVIDENCE_PREVIEW_LIMIT for candidate in first["review_queue_candidates"])


def test_r7cd_load_and_reconcile_reads_xlsx_path(tmp_path: Path) -> None:
    workbook_path = tmp_path / "datefac_raw_material_anjing_foods.xlsx"
    _workbook().save(workbook_path)

    result = load_and_reconcile_real_artifacts(FIXTURE_PATH, workbook_path)

    assert result["summary"]["comparison_row_count"] >= 15
    assert result["summary"]["review_required_count"] == 0


def test_r7cd_module_has_no_db_network_llm_ocr_mineru_imports() -> None:
    source = MODULE_PATH.read_text(encoding="utf-8")
    parsed = ast.parse(source)
    imported_names = {
        alias.name
        for node in ast.walk(parsed)
        if isinstance(node, ast.Import)
        for alias in node.names
    }
    imported_from = {
        node.module
        for node in ast.walk(parsed)
        if isinstance(node, ast.ImportFrom) and node.module is not None
    }
    forbidden_imports = {
        "sqlite3",
        "sqlalchemy",
        "psycopg",
        "psycopg2",
        "pymysql",
        "mysql",
        "asyncpg",
        "socket",
        "subprocess",
        "requests",
        "httpx",
        "openai",
        "anthropic",
        "pytesseract",
        "easyocr",
        "paddleocr",
        "mineru",
    }

    assert imported_names.isdisjoint(forbidden_imports)
    assert all(module not in forbidden_imports for module in imported_from)
    assert "client_ready = true" not in source.lower()
    assert "production_ready = true" not in source.lower()
